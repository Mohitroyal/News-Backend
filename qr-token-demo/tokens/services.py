"""Local demo: one preserved template workbook per token, serialized file writes."""
import hashlib
import json
import os
import re
import tempfile
from pathlib import Path
from uuid import UUID

import qrcode
from django.conf import settings
from filelock import FileLock
from openpyxl import load_workbook

from .models import Token


class InputError(ValueError):
    pass


class ConflictError(ValueError):
    pass


def lock():
    return FileLock(str(settings.DATA_ROOT / 'files.lock'), timeout=15)


def paths(token):
    folder = settings.DATA_ROOT / 'tokens' / str(token.id)
    return folder / 'workbook.xlsx', folder / 'qr.png'


def validate(data):
    if not isinstance(data, dict):
        raise InputError('Send a JSON object.')
    result = {}
    for field, limit in [('token_number', 80), ('project', 160), ('facility', 160)]:
        value = data.get(field)
        if not isinstance(value, str) or not value.strip() or len(value.strip()) > limit:
            raise InputError(f'{field} is required and must be at most {limit} characters.')
        result[field] = value.strip()
    if not re.fullmatch(r'[A-Za-z0-9_-]+', result['token_number']):
        raise InputError('Token number can contain letters, numbers, hyphens and underscores.')
    try:
        request_id = UUID(str(data.get('request_id', '')))
    except ValueError as exc:
        raise InputError('request_id must be a UUID, reused when retrying the same request.') from exc
    return result, request_id


def literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'  # User text must never become an Excel formula.


def field_row(sheet, label):
    for row in sheet.iter_rows(min_row=3):
        if row[0].value == label:
            return row
    raise InputError(f'Missing template field: {label}')


def atomic_save(workbook, target):
    fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=target.parent)
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, target)
    finally:
        Path(temporary).unlink(missing_ok=True)


def write_workbook(token):
    target, _ = paths(token)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        return
    workbook = load_workbook(settings.TEMPLATE_WORKBOOK)
    try:
        summary = workbook['Token Summary']
        old_id = field_row(summary, 'Token ID')[1].value
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str) and old_id in cell.value:
                        literal(cell, cell.value.replace(old_id, token.token_number))
        for label, value in [('Project', token.inputs['project']),
                             ('Facility', token.inputs['facility']),
                             ('Verification status', 'DRAFT - DEMO')]:
            literal(field_row(summary, label)[1], value)
        literal(field_row(summary, 'Token ID')[2], 'DRAFT - DEMO')
        literal(field_row(summary, 'Token ID')[5],
                'Demo copy; synthetic sample data.')
        literal(field_row(summary, 'Verification status')[5],
                'Demo status; not real verification.')
        atomic_save(workbook, target)
    finally:
        workbook.close()


def write_qr(token):
    _, target = paths(token)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        return
    fd, temporary = tempfile.mkstemp(suffix='.png', dir=target.parent)
    os.close(fd)
    try:
        qr = qrcode.QRCode(box_size=8, border=4)
        qr.add_data(token.qr_url)
        qr.make(fit=True)
        qr.make_image(fill_color='black', back_color='white').save(temporary)
        os.replace(temporary, target)
    finally:
        Path(temporary).unlink(missing_ok=True)


def create_token(data):
    inputs, request_id = validate(data)
    digest = hashlib.sha256(json.dumps(inputs, sort_keys=True).encode()).hexdigest()
    with lock():
        token = Token.objects.filter(request_id=request_id).first()
        replay = bool(token and token.state == 'READY')
        if token and token.payload_hash != digest:
            raise ConflictError('This request ID was already used for different input.')
        if not token:
            if Token.objects.filter(token_number=inputs['token_number']).exists():
                raise ConflictError('That token number already exists. Choose another number.')
            token = Token(request_id=request_id, payload_hash=digest,
                          token_number=inputs['token_number'], inputs=inputs)
            # The QR destination is the workbook itself. Scanning it opens/downloads
            # the complete Excel report, without requiring a browser report page.
            token.qr_url = f'{settings.PUBLIC_BASE_URL}/api/tokens/{token.id}/workbook.xlsx'
            token.save()
        if not replay:
            # A failure keeps PENDING; retry resumes this same ID and existing files.
            write_workbook(token)
            write_qr(token)
            token.state = 'READY'
            token.save(update_fields=['state'])
        return token, replay


def metadata(token):
    return {
        'id': str(token.id), 'token_number': token.token_number,
        'state': token.state, 'created_at': token.created_at.isoformat(),
        'qr_url': token.qr_url,
        'detail_url': f'/tokens/{token.id}/',
        'qr_image_url': f'/api/tokens/{token.id}/qr.png',
        'workbook_url': f'/api/tokens/{token.id}/workbook.xlsx',
    }


def read_token(token):
    target, _ = paths(token)
    with lock():
        workbook = load_workbook(target, data_only=True)
        try:
            sheets = []
            for sheet in workbook:
                rows = [[value.isoformat() if hasattr(value, 'isoformat') else value for value in row]
                        for row in sheet.iter_rows(values_only=True)]
                rows = [row for row in rows if any(v is not None for v in row)]
                sheets.append({'name': sheet.title, 'title': rows[0][0],
                               'headers': rows[1],
                               'rows': [row for row in rows[2:] if any(v is not None for v in row)]})
            summary = {row[0]: row[1] for row in sheets[0]['rows']}
            return {**metadata(token), 'summary': summary, 'sheets': sheets}
        finally:
            workbook.close()


def update_status(token, status):
    if status not in ['DRAFT - DEMO', 'IN REVIEW - DEMO', 'COMPLETE - DEMO']:
        raise InputError('Choose one of the supported demo verification statuses.')
    with lock():
        target, _ = paths(token)
        workbook = load_workbook(target)
        try:
            summary = workbook['Token Summary']
            literal(field_row(summary, 'Verification status')[1], status)
            literal(field_row(summary, 'Token ID')[2], status)
            atomic_save(workbook, target)
        finally:
            workbook.close()
