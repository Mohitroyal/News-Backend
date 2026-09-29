import hashlib
import json
import shutil
import tempfile
import uuid
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.test import Client, TestCase, override_settings
from openpyxl import load_workbook

from .models import Token
from . import services


class TokenFlowTests(TestCase):
    def setUp(self):
        self.directory = settings.BASE_DIR / 'data' / 'tests' / str(uuid.uuid4())
        self.directory.mkdir(parents=True)
        self.override = override_settings(DATA_ROOT=self.directory)
        self.override.enable()
        self.addCleanup(shutil.rmtree, self.directory)
        self.addCleanup(self.override.disable)
        self.payload = {'request_id': str(uuid.uuid4()), 'token_number':'TEST-001',
                        'project':'Demo project', 'facility':'Demo facility'}

    def create(self, payload=None):
        return self.client.post('/api/tokens/', json.dumps(payload or self.payload), content_type='application/json')

    def test_complete_create_preserves_workbook_and_serves_files(self):
        result = self.create()
        self.assertEqual(result.status_code, 201)
        data = result.json()
        self.assertEqual(len(data['sheets']), 10)
        self.assertEqual(data['sheets'][0]['headers'][0], 'Field')
        self.assertEqual(data['summary']['Project'], 'Demo project')
        self.assertEqual(data['summary']['Net token removal'], 152.0328)
        token = Token.objects.get()
        original = load_workbook(settings.TEMPLATE_WORKBOOK)
        generated = load_workbook(services.paths(token)[0])
        self.assertEqual(original.sheetnames, generated.sheetnames)
        for before, after in zip(original, generated):
            self.assertEqual(list(before.merged_cells.ranges),list(after.merged_cells.ranges))
            for row in before:
                for cell in row:
                    self.assertEqual(cell.style_id,after[cell.coordinate].style_id)
        original.close(); generated.close()
        for endpoint in [data['detail_url'],data['qr_image_url'],data['workbook_url']]:
            response = self.client.get(endpoint)
            self.assertEqual(response.status_code, 200)
            response.close()
        report = self.client.get(data['detail_url'])
        self.assertContains(report, 'CARBONTRACER TOKEN REPORT')
        self.assertContains(report, 'report-section', count=10)
        self.assertContains(report, 'Sri Lakshmi Coconut Farm')
        self.assertNotContains(report, 'Download Excel')
        self.assertNotContains(report, 'Save to Excel')
        self.assertEqual(self.client.get(data['detail_url'] + 'manage/').status_code, 200)
        self.assertTrue(data['qr_url'].endswith('/workbook.xlsx'))

    def test_retry_and_conflicting_payload(self):
        self.assertEqual(self.create().status_code, 201)
        self.assertEqual(self.create().status_code, 200)
        changed = {**self.payload, 'project':'Changed'}
        self.assertEqual(self.create(changed).status_code,409)
        changed = {**self.payload, 'request_id':str(uuid.uuid4())}
        self.assertEqual(self.create(changed).status_code,409)
        self.assertEqual(Token.objects.count(),1)

    def test_failure_resumes_same_id_and_workbook(self):
        with patch('tokens.services.write_qr', side_effect=OSError('Disk full')):
            self.assertEqual(self.create().status_code,503)
        token = Token.objects.get()
        self.assertEqual(token.state,'PENDING')
        workbook = services.paths(token)[0].read_bytes()
        self.assertEqual(self.client.get(f'/api/tokens/{token.id}/').status_code,404)
        self.assertEqual(self.create().status_code,201)
        self.assertEqual(Token.objects.get().id,token.id)
        self.assertEqual(services.paths(token)[0].read_bytes(),workbook)

    def test_excel_update_changes_lookup_without_changing_qr(self):
        self.create()
        token = Token.objects.get()
        file, qr = services.paths(token)
        before = hashlib.sha256(qr.read_bytes()).hexdigest()
        endpoint = f'/api/tokens/{token.id}/'
        result = self.client.patch(endpoint,json.dumps({'verification_status':'COMPLETE - DEMO'}),content_type='application/json')
        self.assertEqual(result.status_code,200)
        self.assertEqual(result.json()['summary']['Verification status'],'COMPLETE - DEMO')
        # An external edit to the saved Excel file is visible on the next lookup.
        workbook = load_workbook(file)
        services.field_row(workbook['Token Summary'],'Project')[1].value = 'Edited in Excel'
        workbook.save(file); workbook.close()
        self.assertEqual(self.client.get(endpoint).json()['summary']['Project'],'Edited in Excel')
        report = self.client.get(f'/tokens/{token.id}/')
        self.assertContains(report, 'Edited in Excel')
        self.assertContains(report, 'COMPLETE - DEMO')
        self.assertEqual(hashlib.sha256(qr.read_bytes()).hexdigest(),before)
        self.assertIn('no-store',self.client.get(endpoint)['Cache-Control'])

    def test_separate_tokens_have_separate_workbooks(self):
        first = self.create().json()
        second = self.create({**self.payload,'token_number':'TEST-002','request_id':str(uuid.uuid4())}).json()
        self.assertNotEqual(first['id'],second['id'])
        self.client.patch(f"/api/tokens/{first['id']}/",json.dumps({'verification_status':'COMPLETE - DEMO'}),content_type='application/json')
        self.assertEqual(self.client.get(f"/api/tokens/{second['id']}/").json()['summary']['Verification status'],'DRAFT - DEMO')

    def test_validation_csrf_and_formula_injection(self):
        self.assertEqual(self.create({**self.payload,'token_number':'../bad'}).status_code,400)
        self.assertEqual(self.client.post('/api/tokens/','[]',content_type='application/json').status_code,400)
        self.assertEqual(self.client.post('/api/tokens/','{',content_type='application/json').status_code,400)
        self.assertEqual(Client(enforce_csrf_checks=True).post('/api/tokens/',json.dumps(self.payload),content_type='application/json').status_code,403)
        self.create({**self.payload,'project':'=1+1'})
        workbook=load_workbook(services.paths(Token.objects.get())[0])
        cell=services.field_row(workbook['Token Summary'],'Project')[1]
        self.assertEqual(cell.data_type,'s'); self.assertEqual(cell.value,'=1+1')
        workbook.close()
